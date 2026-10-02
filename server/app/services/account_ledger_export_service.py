from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.schemas.account_ledger_schema import (
    ClientSaleRow,
    SupplierPurchasesLedgerResponse,
    SupplierPurchaseRow,
)

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_PAYMENT_LABELS = {
    "efectivo": "Efectivo",
    "debito": "Débito",
    "credito": "Crédito",
    "cheque": "Cheque",
    "otro": "Otro",
}


def _money(value) -> str:
    value = value if value is not None else Decimal("0")
    return (
        f"{Decimal(value):,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def _payment_label(method: str) -> str:
    return _PAYMENT_LABELS.get(method, method)


def _doc_type_label(doc_type: str | None) -> str:
    if not doc_type:
        return "-"
    if doc_type in ("purchase_invoice", "sales_invoice"):
        return "Remito"
    if doc_type in ("purchase_quote", "sales_quote"):
        return "Presupuesto"
    return str(doc_type)


def _products_text(products) -> str:
    return "\n".join(f"{p.product_name} x{p.quantity}" for p in products) or "-"


def _payments_text(payments) -> str:
    if not payments:
        return "Sin pagos"
    return "\n".join(
        f"{_payment_label(p.method)}: ${_money(p.amount)} ({p.date.strftime('%d/%m/%Y')})"
        for p in payments
    )


def _status_label(status: str) -> str:
    return {"paid": "Pagado", "partial": "Parcial", "pending": "Deuda"}.get(status, status)


# ======================================================================
# Excel
# ======================================================================

def _write_xlsx_sheet(ws, headers: list[str], rows: list[list[str]]) -> None:
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    for row in rows:
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_supplier_purchases_xlsx(data: SupplierPurchasesLedgerResponse) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        "Tipo", "N° Doc.", "Fecha", "Proveedor", "Producto/s", "Cantidad",
        "Formas de pago", "Total", "Deuda", "Estado",
    ]

    for group in data.groups:
        sheet_name = (group.sales_rep.name if group.sales_rep else "Sin vendedor")[:31]
        ws = wb.create_sheet(sheet_name)
        rows = [
            [
                _doc_type_label(getattr(row, "document_type", "purchase_invoice")),
                getattr(row, "document_number", None) or "-",
                row.purchase_date.strftime("%d/%m/%Y"),
                row.supplier_name,
                _products_text(row.products),
                row.total_quantity,
                _payments_text(row.payments),
                f"${_money(row.total_amount)}",
                f"${_money(row.balance)}",
                _status_label(row.payment_status),
            ]
            for row in group.rows
        ]
        _write_xlsx_sheet(ws, headers, rows)

    if not wb.sheetnames:
        wb.create_sheet("Proveedores")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _ledger_rows_xlsx(headers: list[str], rows: list[list[str]], sheet_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    _write_xlsx_sheet(ws, headers, rows)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_client_sales_xlsx(rows: list[ClientSaleRow]) -> bytes:
    headers = [
        "Tipo", "N° Doc.", "Fecha", "Cliente", "Producto/s", "Cantidad",
        "Stock actual", "Formas de pago", "Total", "Deuda", "Estado",
    ]
    data_rows = [
        [
            _doc_type_label(getattr(row, "document_type", "sales_invoice")),
            getattr(row, "document_number", None) or "-",
            row.sale_date.strftime("%d/%m/%Y"),
            row.client_name,
            _products_text(row.products),
            row.total_quantity,
            "\n".join(
                f"{p.product_name}: {p.stock_current if p.stock_current is not None else '-'}"
                for p in row.products
            )
            or "-",
            _payments_text(row.payments),
            f"${_money(row.total_amount)}",
            f"${_money(row.balance)}",
            _status_label(row.payment_status),
        ]
        for row in rows
    ]
    return _ledger_rows_xlsx(headers, data_rows, "Clientes")


# ======================================================================
# PDF
# ======================================================================

_STYLES = getSampleStyleSheet()
_CELL_STYLE = ParagraphStyle("cell", parent=_STYLES["Normal"], fontSize=6.5, leading=8.5)
_HEADER_STYLE = ParagraphStyle(
    "header", parent=_STYLES["Normal"], fontSize=6.5, leading=8.5, textColor=colors.white
)
_TITLE_STYLE = ParagraphStyle("title", parent=_STYLES["Heading2"], fontSize=12)


def _p(text) -> Paragraph:
    return Paragraph(str(text).replace("\n", "<br/>"), _CELL_STYLE)


def _p_header(text) -> Paragraph:
    return Paragraph(str(text), _HEADER_STYLE)


def _build_pdf_document(title: str, sections: list[tuple[str, list[str], list[list]]]) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.0 * cm,
        rightMargin=1.0 * cm,
        topMargin=1.0 * cm,
        bottomMargin=1.0 * cm,
    )

    elements = [Paragraph(title, _TITLE_STYLE), Spacer(1, 0.3 * cm)]

    for subtitle, headers, rows in sections:
        if subtitle:
            elements.append(Paragraph(subtitle, _STYLES["Heading3"]))
            elements.append(Spacer(1, 0.15 * cm))

        table_data = [[_p_header(h) for h in headers]]
        if rows:
            table_data += [[_p(cell) for cell in row] for row in rows]
        else:
            table_data.append(
                [_p("Sin datos") if i == 0 else _p("") for i in range(len(headers))]
            )

        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ]
            )
        )
        elements.append(table)
        elements.append(Spacer(1, 0.5 * cm))

    doc.build(elements)
    return buffer.getvalue()


def build_supplier_purchases_pdf(data: SupplierPurchasesLedgerResponse) -> bytes:
    headers = ["Tipo", "N° Doc.", "Fecha", "Proveedor", "Producto/s", "Cant.", "Formas de pago", "Total", "Deuda", "Estado"]
    sections = []
    for group in data.groups:
        subtitle = group.sales_rep.name if group.sales_rep else "Sin vendedor"
        rows = [
            [
                _doc_type_label(getattr(row, "document_type", "purchase_invoice")),
                getattr(row, "document_number", None) or "-",
                row.purchase_date.strftime("%d/%m/%Y"),
                row.supplier_name,
                _products_text(row.products),
                row.total_quantity,
                _payments_text(row.payments),
                f"${_money(row.total_amount)}",
                f"${_money(row.balance)}",
                _status_label(row.payment_status),
            ]
            for row in group.rows
        ]
        sections.append((subtitle, headers, rows))

    return _build_pdf_document("Cuenta corriente — Proveedores", sections)


def build_client_sales_pdf(rows: list[ClientSaleRow]) -> bytes:
    headers = ["Tipo", "N° Doc.", "Fecha", "Cliente", "Producto/s", "Cant.", "Stock actual", "Formas de pago", "Total", "Deuda", "Estado"]
    table_rows = [
        [
            _doc_type_label(getattr(row, "document_type", "sales_invoice")),
            getattr(row, "document_number", None) or "-",
            row.sale_date.strftime("%d/%m/%Y"),
            row.client_name,
            _products_text(row.products),
            row.total_quantity,
            "\n".join(
                f"{p.product_name}: {p.stock_current if p.stock_current is not None else '-'}"
                for p in row.products
            )
            or "-",
            _payments_text(row.payments),
            f"${_money(row.total_amount)}",
            f"${_money(row.balance)}",
            _status_label(row.payment_status),
        ]
        for row in rows
    ]
    return _build_pdf_document(
        "Cuenta corriente — Clientes", [(None, headers, table_rows)]
    )