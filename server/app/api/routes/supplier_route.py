from fastapi import APIRouter, Depends, File, UploadFile, status, Request, Query
from fastapi.responses import Response
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from app.core.rate_limit import limiter
from app.db.base import get_db, get_db_with_commit
from app.core.dependencies import get_current_active_user
from app.models.sales_rep_model import SalesRep
from app.schemas.supplier_schema import (
    SupplierCreate,
    SupplierResponse,
    SupplierUpdate,
    SupplierListResponse,
)
from app.schemas.supplier_import_schema import (
    SupplierImportCommitRequest,
    SupplierImportPreviewResponse,
    SupplierImportCommitResponse,
)
from app.services.supplier_service import SupplierService
from app.services.supplier_import_service import SupplierImportService

router = APIRouter(prefix="/suppliers", tags=["Suppliers"])

@router.post(
    "/",
    response_model=SupplierResponse,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit("5/minute")
def create_supplier(
    request: Request,
    data: SupplierCreate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.create(data)

@router.get("/", response_model=SupplierListResponse)
@limiter.limit("30/minute")
def get_suppliers(
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    is_active: bool | None = Query(None),
    search: str | None = Query(None),
    db: Session = Depends(get_db),
    current_user: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.get_suppliers(
        page=page,
        page_size=page_size,
        is_active=is_active,
        search=search,
    )

@router.get(
    "/{supplier_id}",
    response_model=SupplierResponse,
    status_code=status.HTTP_200_OK,
)
def get_supplier(
    supplier_id: int,
    db: Session = Depends(get_db),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.get_by_id(supplier_id)

@router.patch(
    "/{supplier_id}",
    response_model=SupplierResponse,
    status_code=status.HTTP_200_OK,
)
@limiter.limit("5/minute")
def update_supplier(
    request: Request,
    supplier_id: int,
    data: SupplierUpdate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.update(supplier_id, data)

@router.get(
    "/import/sample-excel",
    summary="Descargar Excel de ejemplo para importar proveedores",
    response_class=Response,
    responses={
        200: {
            "content": {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}
            },
            "description": "Plantilla .xlsx lista para completar y subir a /import/preview.",
        }
    },
)
def download_supplier_import_sample(
    _: SalesRep = Depends(get_current_active_user),
):
    """
    Devuelve un archivo Excel de ejemplo explicando cómo preparar los datos
    para importar proveedores via **POST /api/suppliers/import/preview**.

    La hoja **'Proveedores'** incluye:
    - Encabezados coloreados (azul oscuro = obligatorio, azul claro = opcional).
    - Filas de ejemplo con datos reales para entender el formato esperado.
    - Una hoja **'Instrucciones'** con el paso a paso completo del proceso.
    """
    # ── Estilos ──────────────────────────────────────────────────────────────
    required_fill = PatternFill("solid", fgColor="2E75B6")
    optional_fill = PatternFill("solid", fgColor="70A9D6")
    header_font   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    data_font     = Font(name="Arial", size=10)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    red_font    = Font(name="Arial", bold=True, color="C00000", size=10)
    title_font  = Font(name="Arial", bold=True, size=13)
    step_font   = Font(name="Arial", bold=True, size=11)
    normal_font = Font(name="Arial", size=10)

    wb = openpyxl.Workbook()

    # ── Hoja 1: Datos ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Proveedores"
    ws.freeze_panes = "A2"

    columns = [
        ("name",    True,  36, "Razón social o nombre del proveedor"),
        ("tax_id",  False, 24, "CUIT. Ej: 30-11223344-5"),
        ("email",   False, 32, "Email de contacto. Ej: ventas@proveedor.com"),
        ("phone",   False, 24, "Teléfono con código de país. Ej: +54 3764 223344"),
        ("address", False, 42, "Dirección completa. Ej: Av. Corrientes 789, Posadas"),
    ]

    for col_idx, (col_name, required, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill      = required_fill if required else optional_fill
        cell.font      = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 42

    samples = [
        ["Distribuidora Norte S.A.", "30-11223344-5", "ventas@distnorte.com",    "+54 3764 223344",  "Av. Corrientes 789, Posadas"],
        ["Importadora Sur",          "27-55667788-3", "compras@imporsur.com.ar", "+54 11 4567-8901", "Belgrano 1200, Buenos Aires"],
        ["Proveedor Sin Datos",       None,            None,                      None,               None],
    ]
    for r_idx, row in enumerate(samples, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font      = data_font
            cell.alignment = Alignment(vertical="center")
            cell.border    = thin_border
        ws.row_dimensions[r_idx].height = 18

    # Leyenda
    legend_row = len(samples) + 3
    ws.cell(row=legend_row, column=1, value="LEYENDA:").font = Font(name="Arial", bold=True, size=10)
    req_cell = ws.cell(row=legend_row+1, column=1, value="  Azul oscuro = Campo OBLIGATORIO (no puede estar vacío)")
    req_cell.fill = required_fill
    req_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+1, start_column=1, end_row=legend_row+1, end_column=3)
    opt_cell = ws.cell(row=legend_row+2, column=1, value="  Azul claro  = Campo OPCIONAL (podés dejarlo vacío)")
    opt_cell.fill = optional_fill
    opt_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+2, start_column=1, end_row=legend_row+2, end_column=3)

    # Advertencia
    warn_row = legend_row + 4
    warn_cell = ws.cell(row=warn_row, column=1,
        value="⚠  No elimines ni renombres los encabezados de la fila 1. "
              "No agregues columnas extra. Eliminá estas filas de ejemplo antes de subir el archivo real.")
    warn_cell.font = red_font
    ws.merge_cells(start_row=warn_row, start_column=1, end_row=warn_row, end_column=len(columns))

    # ── Hoja 2: Instrucciones ────────────────────────────────────────────────
    wi = wb.create_sheet("Instrucciones")
    wi.column_dimensions["A"].width = 4
    wi.column_dimensions["B"].width = 90

    r = 1
    cell = wi.cell(row=r, column=1, value="📋  Cómo importar proveedores correctamente")
    cell.font = title_font
    wi.row_dimensions[r].height = 30
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 2

    steps = [
        ("Paso 1 — Completá los datos en la hoja 'Proveedores'",
         [
             "• Completá una fila por proveedor.",
             "• El único campo OBLIGATORIO es 'name' (razón social del proveedor).",
             "• Los demás campos son opcionales: podés dejarlos vacíos.",
             "• Eliminá las filas de ejemplo antes de subir el archivo.",
         ]),
        ("Paso 2 — Guardá el archivo como .xlsx",
         [
             "• Archivo > Guardar como > Formato: Excel (.xlsx).",
             "• No lo guardes como .csv ni .xls, solo .xlsx es aceptado.",
         ]),
        ("Paso 3 — Subí el archivo a 'Vista Previa'",
         [
             "• Usá el endpoint POST /api/suppliers/import/preview enviando el archivo.",
             "• El sistema te mostrará qué filas son válidas y cuáles tienen errores.",
             "• Las filas con errores NO se importan; podés corregirlas y volver a subir.",
         ]),
        ("Paso 4 — Confirmá la importación",
         [
             "• Si el preview es correcto, usá POST /api/suppliers/import/commit.",
             "• El modo por defecto es 'upsert': si el proveedor ya existe (por nombre o tax_id),",
             "  lo actualiza; si no existe, lo crea.",
         ]),
    ]

    for title, bullets in steps:
        title_cell = wi.cell(row=r, column=1, value=title)
        title_cell.font = step_font
        wi.row_dimensions[r].height = 22
        wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for bullet in bullets:
            bc = wi.cell(row=r, column=2, value=bullet)
            bc.font = normal_font
            wi.row_dimensions[r].height = 16
            r += 1
        r += 1

    r += 1
    hdr = wi.cell(row=r, column=1, value="Reglas importantes por columna:")
    hdr.font = Font(name="Arial", bold=True, size=11)
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1

    rules = [
        ("name",    "OBLIGATORIO. Razón social o nombre del proveedor. Mínimo 2 caracteres."),
        ("tax_id",  "Opcional. CUIT u otro identificador fiscal. Se normaliza automáticamente. Ej: 30-11223344-5"),
        ("email",   "Opcional. Dirección de email válida. Ej: ventas@proveedor.com"),
        ("phone",   "Opcional. Teléfono con código de país. Ej: +54 3764 223344"),
        ("address", "Opcional. Dirección completa del proveedor."),
    ]
    for col_name, rule in rules:
        wi.cell(row=r, column=1, value=col_name).font = Font(name="Arial", bold=True, size=10)
        wi.cell(row=r, column=2, value=rule).font = normal_font
        wi.row_dimensions[r].height = 16
        r += 1

    # ── Serializar ───────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=proveedores_importacion_ejemplo.xlsx"},
    )

@router.post(
    "/import/preview",
    response_model=SupplierImportPreviewResponse,
)
@limiter.limit("5/minute")
async def preview_supplier_import(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierImportService(db)

    content = await file.read()

    return service.preview_import(
        content=content,
        filename=file.filename,
    )
    
@router.post("/import/commit", response_model=SupplierImportCommitResponse)
@limiter.limit("5/minute")
async def commit_supplier_import(
    request: Request,
    data: SupplierImportCommitRequest,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """
    Persiste los proveedores validados previamente en la base de datos.
    Soporta modo 'upsert' (crea o actualiza por nombre/tax_id).
    """
    import_service = SupplierImportService(db)
    return import_service.commit_import(data)

@router.post(
    "/import/commit-file",
    response_model=SupplierImportCommitResponse,
)
@limiter.limit("5/minute")
async def commit_supplier_import_file(
    request: Request,
    file: UploadFile = File(...),
    mode: str = Query(default="upsert", pattern="^(create|upsert)$"),
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    """
    Atajo que combina preview + commit en un solo paso.
    Usar /import/preview + /import/commit para mayor control sobre los datos.
    """
    service = SupplierImportService(db)

    content = await file.read()

    preview = service.preview_import(
        content=content,
        filename=file.filename,
    )

    return service.commit_import(
        SupplierImportCommitRequest(
            rows=preview.rows_valid,
            mode=mode,
        )
    )

@router.patch(
    "/{supplier_id}/deactivate",
    response_model=SupplierResponse,
    status_code=status.HTTP_200_OK,
)
def deactivate_supplier(
    supplier_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.deactivate(supplier_id)

@router.patch(
    "/{supplier_id}/activate",
    response_model=SupplierResponse,
    status_code=status.HTTP_200_OK,
)
def activate_supplier(
    supplier_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = SupplierService(db)
    return service.activate(supplier_id)