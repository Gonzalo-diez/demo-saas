from fastapi import APIRouter, Depends, File, UploadFile, Form, Query, Request, status
from fastapi.responses import Response
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from app.db.base import get_db, get_db_with_commit
from app.core.rate_limit import limiter
from app.core.dependencies import get_current_active_user, get_optional_catalog_viewer
from app.models.client_model import Client
from app.models.sales_rep_model import SalesRep
from app.services.product_service import ProductService
from app.services.product_import_service import ProductImportService
from app.services.supplier_pdf_import_service import build_excel
from app.schemas.product_schema import (
    ProductCreateAdmin,
    ProductResponse, 
    ProductUpdate, 
    ProductListResponse,
    ProductFiltersResponse,
    ProductSort,
    PublicProductListResponse,
    PublicProductResponse,
)
from app.schemas.product_import_schema import (
    ProductImportCommitRequest,
    ProductImportCommitResponse,
    ProductImportPreviewResponse,
)

router = APIRouter(prefix="/products", tags=["Products"])

# --- ENDPOINTS DE CONSULTA Y CRUD ---

# Staff -> respuesta completa (con costo, stock mínimo, estado). Visitantes y clientes ->
# vista de tienda, sin datos internos. IMPORTANTE: el modelo completo va primero en la
# unión para que la respuesta del personal no se recorte.
@router.get("/", response_model=ProductListResponse | PublicProductListResponse)
@limiter.limit("60/minute")
def get_products(
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(24, ge=1, le=100),
    search: str | None = Query(None, max_length=150),
    brand: str | None = Query(None),
    category: str | None = Query(None),
    is_active: bool | None = Query(True),
    sort: ProductSort | None = Query(None),
    catalog_only: bool | None = Query(None),
    category_id: int | None = Query(None, gt=0),
    is_public: bool | None = Query(None),
    db: Session = Depends(get_db),
    viewer: SalesRep | Client | None = Depends(get_optional_catalog_viewer),
):
    """Lista productos con filtros normalizados. Visitantes sin login también pueden verla."""

    # Un visitante o un Client (tienda online) SIEMPRE ve solo lo publicado (producto
    # público, activo y de categoría pública): los query `catalog_only`, `is_active` e
    # `is_public` no pueden usarse para saltearlo. Un SalesRep ve todo por defecto;
    # catalog_only=true le sirve para previsualizar lo que ven los clientes.
    is_staff = isinstance(viewer, SalesRep)
    if not is_staff:
        is_catalog_only = True
        is_active = True
        is_public = None
    else:
        is_catalog_only = bool(catalog_only)

    service = ProductService(db)
    result = service.list_products(
        search=search,
        brand=brand,
        category=category,
        is_active=is_active,
        page=page,
        page_size=page_size,
        sort=sort,
        catalog_only=is_catalog_only,
        category_id=category_id,
        is_public=is_public,
    )

    if is_staff:
        return result
    return PublicProductListResponse(
        items=[PublicProductResponse.model_validate(item.model_dump()) for item in result.items],
        total=result.total,
        page=result.page,
        page_size=result.page_size,
    )

@router.get("/filters", response_model=ProductFiltersResponse)
def get_product_filters(
    db: Session = Depends(get_db),
    viewer: SalesRep | Client | None = Depends(get_optional_catalog_viewer),
):
    """Obtiene marcas y categorías únicas para los filtros de la UI."""
    service = ProductService(db)
    return service.get_product_filters(catalog_only=not isinstance(viewer, SalesRep))

# Las categorías se administran en /api/categories (cada distribuidora crea las suyas).

@router.get("/{product_id}", response_model=ProductResponse | PublicProductResponse)
def get_product(
    product_id: int,
    db: Session = Depends(get_db),
    viewer: SalesRep | Client | None = Depends(get_optional_catalog_viewer),
):
    service = ProductService(db)
    is_staff = isinstance(viewer, SalesRep)
    # Un visitante o cliente solo puede ver un producto si está publicado en el catálogo.
    product = service.get_product(product_id, catalog_only=not is_staff)
    if is_staff:
        return product
    return PublicProductResponse.model_validate(product)

@router.post("/", response_model=ProductResponse, status_code=status.HTTP_201_CREATED)
def create_product(
    product_in: ProductCreateAdmin, 
    db: Session = Depends(get_db_with_commit),
    current_user: SalesRep = Depends(get_current_active_user)
):
    """
    Crea un producto aplicando normalización y generación de slug única.

    El precio de venta se manda directo (`unit_price`) o se calcula con `markup_percent`
    sobre el costo (costo 200 + 40% = 280, redondeado al peso entero). El stock inicial
    queda como primera entrada del historial de compras, con su `expiry_date` si se carga.
    """
    service = ProductService(db)
    return service.create_product(product_in=product_in, created_by=current_user.id)

@router.patch("/{product_id}", response_model=ProductResponse)
def update_product(
    product_id: int, 
    product_in: ProductUpdate, 
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    service = ProductService(db)
    return service.update_product(product_id=product_id, obj_in=product_in)

@router.patch("/{product_id}/deactivate", response_model=ProductResponse)
def deactivate_product(
    product_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """Desactivación lógica (Soft Delete) del producto."""
    service = ProductService(db)
    return service.deactivate_product(product_id=product_id)

@router.patch("/{product_id}/reactivate", response_model=ProductResponse)
def reactivate_product(
    product_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """Reactivación de un producto previamente desactivado."""
    service = ProductService(db)
    return service.reactivate_product(product_id=product_id)

@router.patch("/{product_id}/publish", response_model=ProductResponse)
def publish_product(
    product_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """Publica el producto en el catálogo de clientes."""
    service = ProductService(db)
    return service.set_public(product_id, True)

@router.patch("/{product_id}/unpublish", response_model=ProductResponse)
def unpublish_product(
    product_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """Lo saca del catálogo de clientes (sigue disponible para venta B2B)."""
    service = ProductService(db)
    return service.set_public(product_id, False)

@router.post("/import/preview", response_model=ProductImportPreviewResponse)
@limiter.limit("5/minute")
async def preview_product_import(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """Analiza un Excel y devuelve una vista previa de filas válidas y errores."""
    service = ProductImportService(db)
    content = await file.read()
    return service.preview(content)

@router.post("/import/commit", response_model=ProductImportCommitResponse)
@limiter.limit("5/minute")
async def commit_product_import(
    request: Request,
    data: ProductImportCommitRequest,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    """
    Persiste los productos validados previamente en la base de datos.
    Soporta modo 'upsert' (crea o actualiza por SKU).
    """
    service = ProductImportService(db)
    return service.commit(data)

@router.post("/import/commit-file", response_model=ProductImportCommitResponse)
@limiter.limit("5/minute")
async def commit_product_import_from_file(
    request: Request,
    file: UploadFile = File(...),
    mode: str = Form("upsert"),
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user)
):
    service = ProductImportService(db)
    content = await file.read()
    
    # 1. Convertir bytes a filas usando el preview que ya tienes
    preview = service.preview(content)
    
    # 2. Crear el objeto de request manualmente
    commit_request = ProductImportCommitRequest(rows=preview.rows_valid, mode=mode)
    
    # 3. Ejecutar
    return service.commit(commit_request)

@router.get(
    "/import/sample-excel",
    summary="Descargar Excel de ejemplo para importar productos",
    response_class=Response,
    responses={
        200: {
            "content": {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}
            },
            "description": "Plantilla .xlsx lista para completar y subir a /import/preview.",
        }
    },
)
def download_product_import_sample(
    _: SalesRep = Depends(get_current_active_user),
):
    """
    Devuelve un archivo Excel de ejemplo explicando cómo preparar los datos
    para importar productos via **POST /api/products/import/preview**.

    La hoja **'Productos'** incluye:
    - Encabezados coloreados (azul oscuro = obligatorio, azul claro = opcional).
    - Filas de ejemplo con datos reales para entender el formato esperado.
    - Una hoja **'Instrucciones'** con el paso a paso completo del proceso.
    """
    # ── Estilos ──────────────────────────────────────────────────────────────
    required_fill = PatternFill("solid", fgColor="2E75B6")
    optional_fill = PatternFill("solid", fgColor="70A9D6")
    header_font   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    data_font     = Font(name="Arial", size=10)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    red_font    = Font(name="Arial", bold=True, color="C00000", size=10)
    note_font   = Font(name="Arial", italic=True, color="595959", size=9)
    title_font  = Font(name="Arial", bold=True, size=13)
    step_font   = Font(name="Arial", bold=True, size=11)
    normal_font = Font(name="Arial", size=10)

    wb = openpyxl.Workbook()

    # ── Hoja 1: Datos ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Productos"
    ws.freeze_panes = "A2"

    columns = [
        # (nombre_columna, obligatorio, ancho, descripcion)
        ("sku",           True,  16, "Código único del producto. Ej: SKU-001"),
        ("name",          True,  32, "Nombre completo del producto"),
        ("brand",         True,  22, "Marca. Ej: Ledesma"),
        ("category",      True,  22, "Categoría. Ej: Azúcares"),
        ("unit_cost",     True,  16, "Precio de costo (número, sin $). Ej: 420.00"),
        ("unit_price",    True,  16, "Precio de venta (número, sin $). Ej: 650.00"),
        ("stock_current", False, 17, "Stock actual (entero, default 0)"),
        ("stock_min",     False, 14, "Stock mínimo (entero, default 0)"),
        ("description",   False, 42, "Descripción opcional del producto"),
        ("image_url",     False, 42, "URL de imagen (opcional)"),
        ("is_active",     False, 13, "TRUE o FALSE (default TRUE)"),
    ]

    for col_idx, (col_name, required, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill      = required_fill if required else optional_fill
        cell.font      = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 42

    samples = [
        ["SKU-001", "Aceite de Oliva Extra Virgen 500ml", "Cocinero", "Aceites",  850.00, 1200.00, 50,  10, "Aceite importado, botella de vidrio", "https://mi-sitio.com/img/aceite.jpg", True],
        ["SKU-002", "Azúcar Blanca 1kg",                  "Ledesma",  "Azúcares", 420.00,  650.00, 100, 20, "",                                   "",                                   True],
        ["SKU-003", "Harina 000 1kg",                     "Pureza",   "Harinas",  380.00,  580.00,  80, 15, "Triple cero, ideal para pastelería", "",                                   True],
        ["SKU-004", "Producto Inactivo Ejemplo",           "MarcaX",   "Otros",    100.00,  150.00,   0,  0, "Este producto no se mostrará",      "",                                   False],
    ]
    for r_idx, row in enumerate(samples, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font      = data_font
            cell.alignment = Alignment(vertical="center")
            cell.border    = thin_border
        ws.row_dimensions[r_idx].height = 18

    # Leyenda
    legend_row = len(samples) + 3
    ws.cell(row=legend_row,     column=1, value="LEYENDA:").font = Font(name="Arial", bold=True, size=10)
    req_cell = ws.cell(row=legend_row + 1, column=1, value="  Azul oscuro = Campo OBLIGATORIO (no puede estar vacío)")
    req_cell.fill = required_fill
    req_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+1, start_column=1, end_row=legend_row+1, end_column=4)
    opt_cell = ws.cell(row=legend_row + 2, column=1, value="  Azul claro  = Campo OPCIONAL (podés dejarlo vacío)")
    opt_cell.fill = optional_fill
    opt_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+2, start_column=1, end_row=legend_row+2, end_column=4)

    # Advertencia
    warn_row = legend_row + 4
    warn_cell = ws.cell(row=warn_row, column=1,
        value="⚠  No elimines ni renombres los encabezados de la fila 1. "
              "No agregues columnas extra. Eliminá estas filas de ejemplo antes de subir el archivo real.")
    warn_cell.font = red_font
    ws.merge_cells(start_row=warn_row, start_column=1, end_row=warn_row, end_column=len(columns))

    # ── Hoja 2: Instrucciones ────────────────────────────────────────────────
    wi = wb.create_sheet("Instrucciones")
    wi.column_dimensions["A"].width = 4
    wi.column_dimensions["B"].width = 90

    def add_row(row, col, value, font=None, height=None):
        cell = wi.cell(row=row, column=col, value=value)
        if font:
            cell.font = font
        if height:
            wi.row_dimensions[row].height = height
        return cell

    r = 1
    add_row(r, 1, "📋  Cómo importar productos correctamente", title_font, height=30); r += 2

    steps = [
        ("Paso 1 — Completá los datos en la hoja 'Productos'",
         [
             "• Completá una fila por producto.",
             "• Los campos en azul oscuro son OBLIGATORIOS: sku, name, brand, category, unit_cost, unit_price.",
             "• Los campos en azul claro son opcionales: podés dejarlos vacíos.",
             "• No modifiques los nombres de los encabezados (fila 1).",
             "• Eliminá las filas de ejemplo (filas 2 a 5) antes de subir el archivo.",
         ]),
        ("Paso 2 — Guardá el archivo como .xlsx",
         [
             "• Archivo > Guardar como > Formato: Excel (.xlsx).",
             "• No lo guardes como .csv ni .xls, solo .xlsx es aceptado.",
         ]),
        ("Paso 3 — Subí el archivo a 'Vista Previa'",
         [
             "• Usá el endpoint POST /api/products/import/preview enviando el archivo.",
             "• El sistema te mostrará qué filas son válidas y cuáles tienen errores.",
             "• Las filas con errores NO se importan; podés corregirlas y volver a subir.",
         ]),
        ("Paso 4 — Confirmá la importación",
         [
             "• Si el preview te parece correcto, usá POST /api/products/import/commit.",
             "• El modo por defecto es 'upsert': si el SKU ya existe, actualiza el producto;",
             "  si no existe, lo crea.",
         ]),
    ]

    for title, bullets in steps:
        add_row(r, 1, title, step_font, height=22)
        wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for bullet in bullets:
            add_row(r, 2, bullet, normal_font, height=16)
            r += 1
        r += 1

    r += 1
    add_row(r, 1, "Reglas importantes por columna:", Font(name="Arial", bold=True, size=11))
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1

    rules = [
        ("sku",           "OBLIGATORIO. Texto único, máx. 100 caracteres. Se usa para identificar el producto en upserts."),
        ("name",          "OBLIGATORIO. Nombre del producto, máx. 255 caracteres."),
        ("brand",         "OBLIGATORIO. Marca del producto, máx. 100 caracteres."),
        ("category",      "OBLIGATORIO. Categoría del producto, máx. 100 caracteres. Se normaliza automáticamente."),
        ("unit_cost",     "OBLIGATORIO. Número decimal mayor o igual a 0. Sin símbolo $. Ej: 420.00"),
        ("unit_price",    "OBLIGATORIO. Número decimal mayor o igual a 0. Sin símbolo $. Ej: 650.00"),
        ("stock_current", "Opcional. Número entero ≥ 0. Si se omite, se asume 0."),
        ("stock_min",     "Opcional. Número entero ≥ 0. Si se omite, se asume 0."),
        ("description",   "Opcional. Texto libre, máx. 2000 caracteres."),
        ("image_url",     "Opcional. URL completa de la imagen. Ej: https://misitio.com/img.jpg"),
        ("is_active",     "Opcional. TRUE o FALSE. Si se omite, se asume TRUE."),
    ]
    for col_name, rule in rules:
        bold_cell = wi.cell(row=r, column=1, value=col_name)
        bold_cell.font = Font(name="Arial", bold=True, size=10)
        rule_cell = wi.cell(row=r, column=2, value=rule)
        rule_cell.font = normal_font
        wi.row_dimensions[r].height = 16
        r += 1

    # ── Serializar ───────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=productos_importacion_ejemplo.xlsx"},
    )

@router.post("/import/from-pdf")
@limiter.limit("5/minute")
async def generate_import_excel_from_pdf(
    request: Request,
    file: UploadFile = File(...),
    _: SalesRep = Depends(get_current_active_user),
):
    """
    Recibe un PDF de lista de precios de proveedor y devuelve un .xlsx
    listo para ser completado por el equipo y luego subido a /import/preview.
    Solo incluye productos de las categorías habilitadas en el backend.
    """
    content = await file.read()

    # build_excel acepta bytes directamente (además de path)
    xlsx_bytes = build_excel(content)

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=productos_importacion.xlsx"
        },
    )