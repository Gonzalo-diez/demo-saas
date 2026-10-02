from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status, Request
from fastapi.responses import Response
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from app.core.auth_cookies import clear_auth_cookie, set_auth_cookie
from app.core.config import get_settings
from app.core.rate_limit import limiter
from app.core.dependencies import (
    TenantRef,
    bind_tenant_from_ref,
    get_current_active_client,
    get_current_active_user,
    get_current_superuser,
    get_tenant_ref,
)
from app.db.base import get_db, get_db_with_commit
from app.models.client_model import Client
from app.models.sales_rep_model import SalesRep
from app.schemas.client_branch_schema import (
    ClientBranchCreate,
    ClientBranchResponse,
    ClientBranchUpdate,
)
from app.schemas.client_schema import (
    ClientCreate,
    ClientListResponse,
    ClientLogin,
    ClientMapResponse,
    ClientResponse,
    ClientUpdate,
    MessageResponse,
)
from app.schemas.client_import_schema import (
    ClientImportCommitRequest,
    ClientImportCommitResponse,
    ClientImportPreviewResponse,
)
from app.services.client_service import ClientService
from app.services.client_import_service import ClientImportService

settings = get_settings()

router = APIRouter(prefix="/clients", tags=["Clients"])

@router.post("/login", response_model=ClientResponse)
@limiter.limit("5/minute")
def login_client(
    request: Request,
    data: ClientLogin,
    response: Response,
    db: Session = Depends(get_db_with_commit),
    tenant_ref: TenantRef = Depends(get_tenant_ref),
):
    """Login de un cliente (comercio B2B) para poder ver el catálogo."""
    bind_tenant_from_ref(db, tenant_ref, data.tenant_slug)
    service = ClientService(db)
    client, access_token = service.login(data)
    set_auth_cookie(
        response,
        access_token,
        cookie_name=settings.CLIENT_AUTH_COOKIE_NAME,
    )
    return client

@router.post("/logout", response_model=MessageResponse)
@limiter.limit("5/minute")
def logout_client(
    request: Request,
    response: Response,
):
    clear_auth_cookie(response, cookie_name=settings.CLIENT_AUTH_COOKIE_NAME)
    return {"message": "Sesión cerrada correctamente"}

@router.get("/me", response_model=ClientResponse)
def get_me(current_client: Client = Depends(get_current_active_client)):
    return current_client

@router.get("/me/branches", response_model=list[ClientBranchResponse])
def list_my_branches(
    db: Session = Depends(get_db),
    current_client: Client = Depends(get_current_active_client),
):
    """Sucursales del cliente logueado (para elegir a cuál corresponde un pedido)."""
    service = ClientService(db)
    return service.list_client_branches(current_client.id)

@router.get("", response_model=ClientListResponse)
@limiter.limit("30/minute")
def list_clients(
    request: Request,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=100),
    search: str | None = Query(default=None),
    sales_rep_id: int | None = Query(default=None),
    is_active: bool | None = Query(default=None),
    sort: str | None = Query(default="created_at"),
    db: Session = Depends(get_db),
    _: SalesRep = Depends(get_current_active_user),
):
    service = ClientService(db)
    return service.list_clients(
        page=page,
        page_size=page_size,
        search=search,
        sales_rep_id=sales_rep_id,
        is_active=is_active,
        sort=sort,
    )

@router.get("/map", response_model=ClientMapResponse)
@limiter.limit("30/minute")
def get_clients_for_map(
    request: Request,
    search: str | None = Query(default=None),
    sales_rep_id: int | None = Query(default=None),
    is_active: bool | None = Query(default=None),
    db: Session = Depends(get_db),
    _: SalesRep = Depends(get_current_active_user),
):
    service = ClientService(db)
    return service.get_clients_for_map(
        search=search,
        sales_rep_id=sales_rep_id,
        is_active=is_active,
    )

@router.get(
    "/import/sample-excel",
    summary="Descargar Excel de ejemplo para importar clientes",
    response_class=Response,
    responses={
        200: {
            "content": {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}
            },
            "description": "Plantilla .xlsx lista para completar y subir a /import/preview.",
        }
    },
)
def download_client_import_sample(
    _: SalesRep = Depends(get_current_active_user),
):
    """
    Devuelve un archivo Excel de ejemplo explicando cómo preparar los datos
    para importar clientes via **POST /api/clients/import/preview**.

    La hoja **'Clientes'** incluye:
    - Encabezados coloreados (azul oscuro = obligatorio, azul claro = opcional).
    - Filas de ejemplo con datos reales, incluyendo el caso de cliente con múltiples sucursales.
    - Una hoja **'Instrucciones'** con el paso a paso completo del proceso.
    """
    
    required_fill = PatternFill("solid", fgColor="2E75B6")
    optional_fill = PatternFill("solid", fgColor="70A9D6")
    header_font   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    data_font     = Font(name="Arial", size=10)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    red_font    = Font(name="Arial", bold=True, color="C00000", size=10)
    title_font  = Font(name="Arial", bold=True, size=13)
    step_font   = Font(name="Arial", bold=True, size=11)
    normal_font = Font(name="Arial", size=10)

    wb = openpyxl.Workbook()

    # ── Hoja 1: Datos ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Clientes"
    ws.freeze_panes = "A2"

    columns = [
        ("client_name",      True,  28, "Nombre del cliente (empresa o persona)"),
        ("branch_name",      True,  28, "Nombre de la sucursal. Ej: Casa Central"),
        ("email",            False, 26, "Email de acceso a la tienda online (opcional)"),
        ("password",         False, 20, "Contraseña de acceso (opcional: si se deja vacía, se genera una automáticamente y se informa al finalizar)"),
        ("tax_id",           False, 20, "CUIT/DNI. Ej: 30-12345678-9"),
        ("client_type",      False, 16, "'company' o 'individual' (default: company)"),
        ("address",          False, 32, "Dirección de la sucursal"),
        ("city",             False, 18, "Ciudad"),
        ("contact_name",     False, 24, "Nombre del contacto en esa sucursal"),
        ("contact_phone",    False, 22, "Teléfono. Ej: +54 3764 123456"),
        ("lat",              False, 13, "Latitud GPS (decimal). Ej: -27.3671"),
        ("lng",              False, 13, "Longitud GPS (decimal). Ej: -55.8961"),
        ("reference",        False, 28, "Referencia o nota de ubicación"),
        ("is_main",          False, 11, "TRUE si es la sucursal principal"),
        ("is_active",        False, 12, "TRUE o FALSE (default TRUE)"),
        ("branch_is_active", False, 17, "TRUE o FALSE (default TRUE)"),
    ]

    for col_idx, (col_name, required, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill      = required_fill if required else optional_fill
        cell.font      = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 42

    samples = [
        # Cliente con 2 sucursales → mismo client_name, distintos branch_name
        ["Supermercado El Sol", "Casa Central",   "elsol@ejemplo.com", "",                  "30-12345678-9", "company",    "Av. San Martín 123", "Posadas", "Juan Pérez",  "+54 3764 123456", -27.3671, -55.8961, "Local en galería centro", True,  True, True],
        ["Supermercado El Sol", "Sucursal Norte", "elsol@ejemplo.com", "",                  "30-12345678-9", "company",    "Ruta 12 km 5",       "Posadas", "María Gómez", "+54 3764 654321", -27.3500, -55.8800, "Local en barrio norte",   False, True, True],
        # Cliente individual con una sola sucursal
        ["Almacén La Esquina",  "Única",          "laesquina@ejemplo.com", "unaClave123",  "20-98765432-1", "individual", "Belgrano 456",       "Oberá",   "Carlos L.",   "+54 3755 111222", None,     None,     "",                        True,  True, True],
    ]
    for r_idx, row in enumerate(samples, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font      = data_font
            cell.alignment = Alignment(vertical="center")
            cell.border    = thin_border
        ws.row_dimensions[r_idx].height = 18

    # Nota multisucursal
    note_row = len(samples) + 3
    note = ws.cell(row=note_row, column=1,
        value="💡  Un cliente con varias sucursales: repetí el mismo client_name en filas distintas, "
              "con un branch_name diferente en cada una. El sistema las agrupa automáticamente.")
    note.font = Font(name="Arial", italic=True, color="1F4E79", size=9, bold=True)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)

    # Nota sobre contraseña
    password_note_row = note_row + 1
    password_note = ws.cell(row=password_note_row, column=1,
        value="💡  Si dejás 'password' vacío al crear un cliente nuevo, el sistema genera una "
              "contraseña automática y te la muestra al terminar de importar, para que se la "
              "pases al cliente. Si estás actualizando un cliente que ya existe, dejar 'password' "
              "vacío NO le cambia la contraseña.")
    password_note.font = Font(name="Arial", italic=True, color="1F4E79", size=9, bold=True)
    ws.merge_cells(start_row=password_note_row, start_column=1, end_row=password_note_row, end_column=8)

    # Leyenda
    legend_row = password_note_row + 2
    ws.cell(row=legend_row, column=1, value="LEYENDA:").font = Font(name="Arial", bold=True, size=10)
    req_cell = ws.cell(row=legend_row+1, column=1, value="  Azul oscuro = Campo OBLIGATORIO (no puede estar vacío)")
    req_cell.fill = required_fill
    req_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+1, start_column=1, end_row=legend_row+1, end_column=5)
    opt_cell = ws.cell(row=legend_row+2, column=1, value="  Azul claro  = Campo OPCIONAL (podés dejarlo vacío)")
    opt_cell.fill = optional_fill
    opt_cell.font = Font(name="Arial", color="FFFFFF", size=9)
    ws.merge_cells(start_row=legend_row+2, start_column=1, end_row=legend_row+2, end_column=5)

    # Advertencia
    warn_row = legend_row + 4
    warn_cell = ws.cell(row=warn_row, column=1,
        value="⚠  No elimines ni renombres los encabezados de la fila 1. "
              "No agregues columnas extra. Eliminá estas filas de ejemplo antes de subir el archivo real.")
    warn_cell.font = red_font
    ws.merge_cells(start_row=warn_row, start_column=1, end_row=warn_row, end_column=len(columns))

    # ── Hoja 2: Instrucciones ────────────────────────────────────────────────
    wi = wb.create_sheet("Instrucciones")
    wi.column_dimensions["A"].width = 4
    wi.column_dimensions["B"].width = 90

    r = 1
    cell = wi.cell(row=r, column=1, value="📋  Cómo importar clientes correctamente")
    cell.font = title_font
    wi.row_dimensions[r].height = 30
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 2

    steps = [
        ("Paso 1 — Completá los datos en la hoja 'Clientes'",
         [
             "• Completá una fila por SUCURSAL (no por cliente).",
             "• Un cliente con varias sucursales ocupa varias filas: repetí el client_name igual en cada una.",
             "• Los campos en azul oscuro son OBLIGATORIOS: client_name y branch_name.",
             "• Los campos en azul claro son opcionales.",
             "• lat y lng deben venir siempre juntos (o ambos vacíos).",
             "• Eliminá las filas de ejemplo antes de subir el archivo.",
         ]),
        ("Paso 2 — Guardá el archivo como .xlsx",
         [
             "• Archivo > Guardar como > Formato: Excel (.xlsx).",
             "• No lo guardes como .csv ni .xls, solo .xlsx es aceptado.",
         ]),
        ("Paso 3 — Subí el archivo a 'Vista Previa'",
         [
             "• Usá el endpoint POST /api/clients/import/preview enviando el archivo.",
             "• El sistema te mostrará qué filas son válidas y cuáles tienen errores.",
             "• Las filas con errores NO se importan; podés corregirlas y volver a subir.",
         ]),
        ("Paso 4 — Confirmá la importación",
         [
             "• Si el preview es correcto, usá POST /api/clients/import/commit.",
             "• El modo por defecto es 'upsert': si el cliente/sucursal ya existe, lo actualiza;",
             "  si no existe, lo crea.",
         ]),
    ]

    for title, bullets in steps:
        title_cell = wi.cell(row=r, column=1, value=title)
        title_cell.font = step_font
        wi.row_dimensions[r].height = 22
        wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for bullet in bullets:
            bc = wi.cell(row=r, column=2, value=bullet)
            bc.font = normal_font
            wi.row_dimensions[r].height = 16
            r += 1
        r += 1

    r += 1
    hdr = wi.cell(row=r, column=1, value="Reglas importantes por columna:")
    hdr.font = Font(name="Arial", bold=True, size=11)
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1

    rules = [
        ("client_name",      "OBLIGATORIO. Nombre del cliente. Mismo valor en todas las sucursales del mismo cliente."),
        ("branch_name",      "OBLIGATORIO. Nombre único de la sucursal dentro del cliente."),
        ("tax_id",           "Opcional. CUIT, DNI u otro identificador fiscal. Ej: 30-12345678-9"),
        ("client_type",      "Opcional. 'company' (empresa) o 'individual' (persona). Default: 'company'."),
        ("address",          "Opcional. Dirección completa de la sucursal. Máx. 255 caracteres."),
        ("city",             "Opcional. Ciudad. Máx. 100 caracteres."),
        ("contact_name",     "Opcional. Nombre del contacto en esa sucursal. Máx. 255 caracteres."),
        ("contact_phone",    "Opcional. Teléfono con código de país. Ej: +54 3764 123456"),
        ("lat",              "Opcional. Latitud en decimal. Ej: -27.3671. Debe venir junto con lng."),
        ("lng",              "Opcional. Longitud en decimal. Ej: -55.8961. Debe venir junto con lat."),
        ("reference",        "Opcional. Nota o referencia de ubicación. Máx. 255 caracteres."),
        ("is_main",          "Opcional. TRUE si es la sucursal principal del cliente. Default: FALSE."),
        ("is_active",        "Opcional. TRUE o FALSE para el cliente. Default: TRUE."),
        ("branch_is_active", "Opcional. TRUE o FALSE para esta sucursal. Default: TRUE."),
    ]
    for col_name, rule in rules:
        wi.cell(row=r, column=1, value=col_name).font = Font(name="Arial", bold=True, size=10)
        wi.cell(row=r, column=2, value=rule).font = normal_font
        wi.row_dimensions[r].height = 16
        r += 1

    # ── Serializar ───────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=clientes_importacion_ejemplo.xlsx"},
    )

@router.post("/import/preview", response_model=ClientImportPreviewResponse)
@limiter.limit("5/minute")
async def preview_client_import(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = ClientImportService(db)
    content = await file.read()
    return service.preview_import(
        content=content,
        filename=file.filename,
    )

@router.post("/import/commit", response_model=ClientImportCommitResponse)
@limiter.limit("5/minute")
async def commit_client_import(
    request: Request,
    body: ClientImportCommitRequest,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_active_user),
):
    service = ClientImportService(db)
    return service.commit_from_rows(
        rows=body.rows,
        mode=body.mode,
    )

@router.post("/import/commit-file", response_model=ClientImportCommitResponse)
@limiter.limit("5/minute")
async def commit_client_file_import(
    request: Request,
    file: UploadFile = File(...),
    mode: str = Query(default="upsert", pattern="^(upsert|create|update)$"),
    db: Session = Depends(get_db_with_commit),
    current_user: SalesRep = Depends(get_current_active_user),
):
    service = ClientImportService(db)
    content = await file.read()
    return service.commit_import(
        content=content,
        filename=file.filename,
        sales_rep_id=current_user.id,
        mode=mode,
    )

@router.get("/{client_id}", response_model=ClientResponse)
def get_client(
    client_id: int,
    db: Session = Depends(get_db),
    _: SalesRep = Depends(get_current_active_user),
):
    try:
        service = ClientService(db)
        return service.get_client(client_id)
    except ValueError as e:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(e),
        )

@router.post("", response_model=ClientResponse, status_code=status.HTTP_201_CREATED)
@limiter.limit("5/minute")
def create_client(
    request: Request,
    data: ClientCreate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.create_client(data)
    except ValueError as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(e),
        )

@router.put("/{client_id}", response_model=ClientResponse)
@limiter.limit("5/minute")
def update_client(
    request: Request,
    client_id: int,
    data: ClientUpdate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.update_client(client_id, data)
    except ValueError as e:
        message = str(e)
        if "no encontrado" in message.lower():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=message,
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=message,
        )

@router.patch("/{client_id}/deactivate", response_model=ClientResponse)
def deactivate_client(
    client_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.deactivate_client(client_id)
    except ValueError as e:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(e),
        )

@router.patch("/{client_id}/activate", response_model=ClientResponse)
def activate_client(
    client_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.activate_client(client_id)
    except ValueError as e:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(e),
        )

@router.get("/{client_id}/branches", response_model=list[ClientBranchResponse])
@limiter.limit("30/minute")
def list_client_branches(
    request: Request,
    client_id: int,
    db: Session = Depends(get_db),
    _: SalesRep = Depends(get_current_active_user),
):
    try:
        service = ClientService(db)
        return service.list_client_branches(client_id)
    except ValueError as e:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(e),
        )

@router.post(
    "/{client_id}/branches",
    response_model=ClientBranchResponse,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit("5/minute")
def create_client_branch(
    request: Request,
    client_id: int,
    data: ClientBranchCreate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.create_branch(client_id, data)
    except ValueError as e:
        message = str(e)
        if "no encontrado" in message.lower():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=message,
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=message,
        )

@router.put("/{client_id}/branches/{branch_id}", response_model=ClientBranchResponse)
@limiter.limit("5/minute")
def update_client_branch(
    request: Request,
    client_id: int,
    branch_id: int,
    data: ClientBranchUpdate,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.update_branch(client_id, branch_id, data)
    except ValueError as e:
        message = str(e)
        if (
            "no encontrado" in message.lower()
            or "no pertenece" in message.lower()
        ):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=message,
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=message,
        )

@router.patch(
    "/{client_id}/branches/{branch_id}/deactivate",
    response_model=ClientBranchResponse,
)
def deactivate_client_branch(
    client_id: int,
    branch_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.deactivate_branch(client_id, branch_id)
    except ValueError as e:
        message = str(e)
        if (
            "no encontrado" in message.lower()
            or "no pertenece" in message.lower()
        ):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=message,
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=message,
        )

@router.patch(
    "/{client_id}/branches/{branch_id}/activate",
    response_model=ClientBranchResponse,
)
def activate_client_branch(
    client_id: int,
    branch_id: int,
    db: Session = Depends(get_db_with_commit),
    _: SalesRep = Depends(get_current_superuser),
):
    try:
        service = ClientService(db)
        return service.activate_branch(client_id, branch_id)
    except ValueError as e:
        message = str(e)
        if (
            "no encontrado" in message.lower()
            or "no pertenece" in message.lower()
        ):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=message,
            )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=message,
        )